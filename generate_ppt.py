"""
PowerPoint Generator for PhishAware Project Presentation.
Builds a 16:9 widescreen presentation deck using python-pptx.
"""

import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


def build_presentation(output_path="PhishAware_Project_Presentation.pptx"):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]  # Blank slide

    # Color Palette
    C_BG = RGBColor(15, 23, 42)          # Deep Navy Slate (#0F172A)
    C_CARD = RGBColor(30, 41, 59)        # Card Navy (#1E293B)
    C_CARD_LIGHT = RGBColor(42, 55, 86)  # Card Border/Surface
    C_ACCENT = RGBColor(59, 130, 246)     # Primary Blue (#3B82F6)
    C_CYAN = RGBColor(56, 189, 248)       # Tech Cyan (#38BDF8)
    C_GREEN = RGBColor(16, 185, 129)      # Emerald Green (#10B981)
    C_AMBER = RGBColor(245, 158, 11)      # Amber (#F59E0B)
    C_RED = RGBColor(239, 68, 68)         # Danger Red (#EF4444)
    C_WHITE = RGBColor(248, 250, 252)     # Off-white (#F8FAFC)
    C_MUTED = RGBColor(148, 163, 184)     # Muted Slate (#94A3B8)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="CYBERSECURITY & DEFENSIVE APPLICATION DESIGN"):
        # Category Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_CYAN

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = C_WHITE

    def add_card(slide, left, top, width, height, title, items, badge_text=None, border_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = border_color if border_color else C_CARD_LIGHT
        card.line.width = Pt(1.5)

        tx_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True

        # Card Title
        p_head = tf.paragraphs[0]
        p_head.text = title
        p_head.font.size = Pt(18)
        p_head.font.bold = True
        p_head.font.color.rgb = C_CYAN
        p_head.space_after = Pt(10)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(13)
            p.font.color.rgb = C_WHITE
            p.space_after = Pt(8)

        return card

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative banner shape
    dec = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.15), Inches(5.0))
    dec.fill.solid()
    dec.fill.fore_color.rgb = C_ACCENT
    dec.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.3), Inches(11), Inches(4.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "DEFENSIVE CYBERSECURITY & SECURE SOFTWARE ENGINEERING"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = C_CYAN
    p_badge.space_after = Pt(14)

    p_main = tf1.add_paragraph()
    p_main.text = "PhishAware: Secure Phishing Awareness\n& URL Risk Checker"
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = C_WHITE
    p_main.space_after = Pt(16)

    p_sub = tf1.add_paragraph()
    p_sub.text = "An educational, defense-in-depth web application evaluating URLs through transparent,\nexplainable heuristics with zero server outbound requests."
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = C_MUTED
    p_sub.space_after = Pt(32)

    p_meta = tf1.add_paragraph()
    p_meta.text = "Tech Stack: Python 3.14  |  Flask 3.x  |  MySQL 8.0 & SQLite Dual-Engine  |  HTML5 & CSS3  |  Git"
    p_meta.font.size = Pt(13)
    p_meta.font.bold = True
    p_meta.font.color.rgb = C_GREEN

    # =========================================================================
    # SLIDE 2: Problem Statement & Motivation
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Problem Statement: The Deception Gap in URL Security")

    add_card(
        s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "The Phishing Threat Landscape",
        [
            "Phishing is the #1 initial access vector in global cyber incidents.",
            "Cybercriminals weaponize RFC 3986 URL parsing rules to visually disguise destinations.",
            "Deceptive tactics include userinfo '@' credential injection, Cyrillic lookalike homoglyphs, and nested subdomains.",
            "Non-technical users cannot decipher complex address structures without guidance."
        ],
        border_color=C_RED
    )

    add_card(
        s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
        "Limitations of Existing Tools",
        [
            "Black-Box Scoring: Most scanners display 'Safe' or 'Dangerous' without explaining the underlying reasoning.",
            "Lack of Education: Users fail to build intuition for spotting phishing patterns independently.",
            "Severe Server Vulnerabilities: Naive scanners fetch or crawl target links, opening servers to SSRF and drive-by malware.",
            "Our Mission: Build an explainable, safe, and educational defensive analyzer."
        ],
        border_color=C_AMBER
    )

    # =========================================================================
    # SLIDE 3: Core Security Philosophy
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Foundational Principles: Zero Outbound Requests & Transparency")

    add_card(
        s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "1. Zero Outbound Requests Protocol",
        [
            "The backend server NEVER crawls, fetches, or resolves DNS for submitted URLs.",
            "Eliminates Server-Side Request Forgery (SSRF) against cloud metadata (169.254.169.254) or local databases.",
            "Prevents Denial-of-Service attacks (Slowloris/tar-pits) caused by malicious remote servers holding connections open.",
            "Protects server memory and disks from automated malware payload downloads."
        ],
        border_color=C_GREEN
    )

    add_card(
        s3, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
        "2. Heuristics are Not Definitive Proof",
        [
            "Heuristic evaluation analyzes syntactic and structural patterns for known deception tactics.",
            "Heuristic results ARE NOT definitive proof that a site is benign or malicious.",
            "Mandatory Disclaimers: Highlight that clean scores do not guarantee safety, and legitimate sites may trigger warnings.",
            "Proactive Education: Teaches users to verify root domains and navigate via bookmarks."
        ],
        border_color=C_CYAN
    )

    # =========================================================================
    # SLIDE 4: System Architecture & Workflow
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "System Architecture: End-to-End Processing Workflow")

    col_w = Inches(2.7)
    gap = Inches(0.25)
    top_y = Inches(1.7)
    h_box = Inches(5.0)

    add_card(
        s4, Inches(0.8), top_y, col_w, h_box,
        "1. Client Input",
        [
            "User submits URL via responsive UI.",
            "Instant synthetic chips provided for live test demos.",
            "Vanilla JS fetch() eliminates CDN dependencies.",
            "Strict textContent DOM rendering (Zero XSS)."
        ]
    )

    add_card(
        s4, Inches(0.8) + (col_w + gap), top_y, col_w, h_box,
        "2. Safe Normalizer",
        [
            "Length cap: Max 2,048 chars prevents ReDoS.",
            "Control chars and null bytes rejected.",
            "Blocks pseudo-schemes (javascript:, data:, file:).",
            "Infers scheme if missing for RFC compliance."
        ]
    )

    add_card(
        s4, Inches(0.8) + (col_w + gap) * 2, top_y, col_w, h_box,
        "3. Heuristic Engine",
        [
            "Offline static parsing via urllib & ipaddress.",
            "Evaluates 13 rules across Host, Authority, Path.",
            "Computes cumulative score (0-100 max cap).",
            "Classifies tier: Low, Suspicious, or High Risk."
        ]
    )

    add_card(
        s4, Inches(0.8) + (col_w + gap) * 3, top_y, col_w, h_box,
        "4. Storage & Output",
        [
            "Dual-Engine persistence (MySQL 8.0 & SQLite).",
            "100% Parameterized queries (%s / ?).",
            "Relational foreign keys with ON DELETE CASCADE.",
            "Returns explainable evaluation report."
        ]
    )

    # =========================================================================
    # SLIDE 5: The 13 Heuristic Rules Catalog
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "The 13 Heuristic Rules: Detectable Deceptive Tactics")

    add_card(
        s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Hostname & Authority Rules",
        [
            "HR-01 (Raw IP Hostname): +35 pts | Disposable phishing kits bypassing domain registration.",
            "HR-02 (Credential '@' Trick): +40 pts | Browser ignores text before '@' and navigates to the host after.",
            "HR-03 (Excessive Subdomains): +20 pts | 3+ subdomains push the real domain offscreen.",
            "HR-04 (Sensitive Subdomain Keyword): +25 pts | 'login' or 'verify' in subdomains of fake domains.",
            "HR-05 (High-Risk TLD): +15 pts | Disposable TLDs (.xyz, .top, .click, .buzz).",
            "HR-06 (Punycode / Homoglyphs): +35 pts | xn-- lookalike Cyrillic characters mimicking brands.",
            "HR-07 (Non-Standard Ports): +20 pts | Web services running outside standard 80/443."
        ]
    )

    add_card(
        s5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
        "Structural, Path & Query Rules",
        [
            "HR-08 (Excessive URL Length): +15 pts | URLs >100 chars hiding redirect tokens or payloads.",
            "HR-09 (Excessive Hyphens): +15 pts | 3+ hyphens creating deceptive imitation domain names.",
            "HR-10 (Dword / Hex IP Obfuscation): +40 pts | Integer/Hex IPs bypassing naive regex filters.",
            "HR-11 (Open Redirect Parameter): +20 pts | Bouncing victims from trusted sites to malicious hosts.",
            "HR-12 (Brand in Path / Query): +25 pts | Brand name in path (e.g. site.com/paypal/login).",
            "HR-13 (Unencrypted HTTP Sensitive): +15 pts | Cleartext HTTP for credential actions."
        ]
    )

    # =========================================================================
    # SLIDE 6: Explainable Scoring & Risk Tiers
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Calibrated Scoring Model & Explainable Feedback")

    c_w = Inches(3.6)
    c_gap = Inches(0.4)

    add_card(
        s6, Inches(0.8), Inches(1.6), c_w, Inches(5.2),
        "🟢 Low Risk (0 - 24 pts)",
        [
            "Score Range: 0 to 24 points.",
            "Syntactic Indicators: Few or no structural anomalies detected.",
            "UI Visual: Emerald Green badge.",
            "Guidance: Benign structure does not guarantee safety; verify the domain ownership independently."
        ],
        border_color=C_GREEN
    )

    add_card(
        s6, Inches(0.8) + c_w + c_gap, Inches(1.6), c_w, Inches(5.2),
        "🟡 Suspicious (25 - 59 pts)",
        [
            "Score Range: 25 to 59 points.",
            "Syntactic Indicators: Multiple warning signs (e.g. high-risk TLD + sensitive keywords + long URL).",
            "UI Visual: Amber warning badge.",
            "Guidance: Exercise high caution; verify URL source before interacting."
        ],
        border_color=C_AMBER
    )

    add_card(
        s6, Inches(0.8) + (c_w + c_gap) * 2, Inches(1.6), c_w, Inches(5.2),
        "🔴 High Risk (60 - 100 pts)",
        [
            "Score Range: 60 to 100 points.",
            "Syntactic Indicators: Critical deception tactics (e.g. '@' symbol, raw IP host, dword IP obfuscation).",
            "UI Visual: Red alert badge.",
            "Guidance: Strong indicators of credential harvesting; do not enter sensitive info."
        ],
        border_color=C_RED
    )

    # =========================================================================
    # SLIDE 7: Database Design (MySQL & Resilient Dual-Engine)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Database Architecture: 3NF Relational Model & SQLi Defense")

    add_card(
        s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Relational Schema (3NF)",
        [
            "Table 1: 'assessments': Stores raw URL, normalized URL, host, scheme, risk score, classification, and timestamp.",
            "Table 2: 'triggered_rules': Stores individual rule activations linked via Foreign Key (assessment_id).",
            "ON DELETE CASCADE preserves referential integrity upon deletion.",
            "Strategic B-Tree Indexes on host, risk_level, created_at, and rule_id for sub-millisecond query response."
        ]
    )

    add_card(
        s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
        "Resilient Dual-Engine & Security",
        [
            "Primary Engine: MySQL 8.0 with PyMySQL client.",
            "Auto-Fallback: Zero-config SQLite (phishing_checker.db) activates if MySQL credentials are not set.",
            "100% Parameterized Statements (%s / ?): Eliminates SQL Injection.",
            "Verified Immunity: Synthetic injection payload (' OR '1'='1) tested and safely treated as plain string data."
        ],
        border_color=C_CYAN
    )

    # =========================================================================
    # SLIDE 8: Web Interface & User Experience
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Web Interface: Clean, Accessible & Educational")

    add_card(
        s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Interactive Analyzer Dashboard",
        [
            "Interactive Test Chips: Instant one-click testing for benign, '@' trick, raw IP, and subdomain spoofing.",
            "Dynamic Risk Gauge: Visual 0-100 score indicator with color transitions.",
            "Deconstructed URL Grid: Protocol scheme, identified host, and normalized components.",
            "Explainable Rule Cards: Detailing 'What Was Found', 'Attacker Tactics', and 'Safety Action'."
        ]
    )

    add_card(
        s8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
        "Audit Trail, Analytics & Quiz",
        [
            "Searchable History: Debounced hostname search with risk-tier filter chips.",
            "Telemetry Dashboard: Metric cards for total scans, average score, and top 5 triggered rules leaderboard.",
            "Interactive Phishing Quiz: Real-world scenario questions with immediate educational feedback.",
            "Zero CDN Dependencies: 100% vanilla JavaScript protects client integrity."
        ]
    )

    # =========================================================================
    # SLIDE 9: Defense-in-Depth & Security Hardening
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Defense-in-Depth: Multi-Layered Security Architecture")

    add_card(
        s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Browser & HTTP Header Hardening",
        [
            "Content-Security-Policy (CSP): Restricts script execution to approved self-hosted origins.",
            "X-Frame-Options: DENY: Blocks clickjacking by preventing iframe embedding.",
            "X-Content-Type-Options: nosniff: Disables MIME-type sniffing.",
            "Referrer-Policy: strict-origin-when-cross-origin: Prevents leakage of URL tokens in referer headers."
        ],
        border_color=C_ACCENT
    )

    add_card(
        s9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
        "Application & Storage Hardening",
        [
            "Contextual HTML Escaping: Jinja2 template auto-escaping prevents Cross-Site Scripting (XSS).",
            "Client-Side DOM Sanitization: textContent used exclusively for rendering user data.",
            "Length & Character Guards: 2,048-char limits block ReDoS and memory bloat.",
            "Centralized Error Handlers: Intercept 400/404/500 errors to prevent stack trace leaks."
        ],
        border_color=C_GREEN
    )

    # =========================================================================
    # SLIDE 10: Verification & Test Results
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Verification & Automated Testing: 100% Pass Rate")

    add_card(
        s10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Comprehensive Pytest Suite (37 Tests)",
        [
            "test_normalizer.py (14 tests): Validated RFC compliance, port boundaries, control char filtering, and blocked pseudo-schemes.",
            "test_heuristics.py (13 tests): Verified individual rule triggers, scoring clamping (0-100), and educational disclaimers.",
            "test_db.py (2 tests): Confirmed relational mapping and verified SQL injection immunity against attack payloads.",
            "test_api.py (8 tests): Verified endpoints, search filters, error responses, and HTTP security headers."
        ]
    )

    add_card(
        s10, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
        "Test Execution Metrics",
        [
            "Platform: Python 3.14 on Windows 64-bit.",
            "Test Framework: pytest 9.2.0 with pluggy.",
            "Total Tests: 37 collected, 37 passed, 0 failures.",
            "Execution Speed: 4.58 seconds for complete end-to-end suite.",
            "Git Status: Clean commit tree tracking each architectural milestone."
        ],
        border_color=C_GREEN
    )

    # =========================================================================
    # SLIDE 11: Limitations & Future Enhancements
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Project Limitations & Future Development Roadmap")

    add_card(
        s11, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Current System Limitations",
        [
            "Static Heuristics Scope: Cannot inspect dynamic DOM content or detect real-time server cloaking.",
            "Zero Network Access: Does not verify live SSL/TLS certificate validity or DNS existence.",
            "False Positives: Complex legitimate developer URLs with multiple subdomains may receive warning flags."
        ],
        border_color=C_AMBER
    )

    add_card(
        s11, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
        "Future Roadmap Enhancements",
        [
            "Offline Machine Learning: Integrate a lightweight Random Forest model trained on URL lexical features.",
            "Threat Feed Bloom Filters: Local, compressed Bloom filters for known malicious host databases (URLhaus).",
            "Browser Extension: Client-side link inspector evaluating URLs in the browser before navigation.",
            "Multi-Language Support: Localization for global educational outreach."
        ],
        border_color=C_CYAN
    )

    # =========================================================================
    # SLIDE 12: Conclusion & Q&A
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)

    dec12 = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.15), Inches(5.0))
    dec12.fill.solid()
    dec12.fill.fore_color.rgb = C_GREEN
    dec12.line.fill.background()

    tb12 = s12.shapes.add_textbox(Inches(1.2), Inches(1.3), Inches(11), Inches(4.8))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    p12_badge = tf12.paragraphs[0]
    p12_badge.text = "PROJECT SUMMARY & VIVA DEFENSE READY"
    p12_badge.font.size = Pt(13)
    p12_badge.font.bold = True
    p12_badge.font.color.rgb = C_CYAN
    p12_badge.space_after = Pt(14)

    p12_main = tf12.add_paragraph()
    p12_main.text = "PhishAware: Empowering Safe Browsing"
    p12_main.font.size = Pt(36)
    p12_main.font.bold = True
    p12_main.font.color.rgb = C_WHITE
    p12_main.space_after = Pt(16)

    p12_sub = tf12.add_paragraph()
    p12_sub.text = "• Proved that explainable heuristic analysis effectively educates users on phishing tactics.\n• Enforced a Zero Outbound Requests policy eliminating SSRF and server infection vectors.\n• Achieved 100% SQL injection immunity and hardened HTTP headers.\n• Full deliverables: Source code, 37 automated tests, SRS, Diagrams, Report & Presentation."
    p12_sub.font.size = Pt(16)
    p12_sub.font.color.rgb = C_WHITE
    p12_sub.space_after = Pt(28)

    p12_qa = tf12.add_paragraph()
    p12_qa.text = "Thank you! Open for Questions & Technical Discussion."
    p12_qa.font.size = Pt(18)
    p12_qa.font.bold = True
    p12_qa.font.color.rgb = C_CYAN

    prs.save(output_path)
    print(f"[+] Successfully generated PowerPoint presentation: {output_path}")


if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "PhishAware_Project_Presentation.pptx"
    build_presentation(out_file)
